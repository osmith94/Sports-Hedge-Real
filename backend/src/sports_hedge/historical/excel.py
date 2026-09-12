from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from sports_hedge.historical.coverage import CoverageReporter
from sports_hedge.historical.repository import SqliteHistoricalRepository

SHEET_NAMES = (
    "Matches",
    "Team Stats",
    "Match Events",
    "Lineups",
    "Coverage",
    "Sources",
)


class HistoricalExcelExporter:
    """Export normalized repository facts to a review workbook."""

    def __init__(self, repository: SqliteHistoricalRepository) -> None:
        self._repository = repository

    def export(self, path: str | Path) -> Path:
        destination = Path(path)
        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)

        matches_sheet = workbook.create_sheet("Matches")
        _write_sheet(
            matches_sheet,
            [
                "match_id",
                "competition",
                "season",
                "kickoff_utc",
                "home_team",
                "away_team",
                "home_ft_goals",
                "away_ft_goals",
                "home_ht_goals",
                "away_ht_goals",
                "status",
                "venue_name",
            ],
            [
                [
                    match.match_id,
                    match.competition_name,
                    match.season_label,
                    match.kickoff_utc.isoformat(),
                    match.home_team_name,
                    match.away_team_name,
                    match.home_ft_goals,
                    match.away_ft_goals,
                    match.home_ht_goals,
                    match.away_ht_goals,
                    match.status.value,
                    match.venue_name,
                ]
                for match in self._repository.list_matches()
            ],
        )

        stats_sheet = workbook.create_sheet("Team Stats")
        _write_sheet(
            stats_sheet,
            [
                "match_id",
                "competition_id",
                "season_id",
                "team",
                "is_home",
                "corners",
                "yellow_cards",
                "red_cards",
                "penalties",
            ],
            [
                [
                    row.match_id,
                    row.competition_id,
                    row.season_id,
                    row.team_name,
                    row.is_home,
                    row.corners,
                    row.yellow_cards,
                    row.red_cards,
                    row.penalties,
                ]
                for row in self._repository.list_team_stats()
            ],
        )

        events_sheet = workbook.create_sheet("Match Events")
        _write_sheet(
            events_sheet,
            [
                "event_id",
                "match_id",
                "event_type",
                "minute",
                "extra_minute",
                "team",
                "player_name",
                "related_player_name",
                "period",
                "source_name",
                "source_event_id",
                "source_observation_id",
            ],
            [
                [
                    event.event_id,
                    event.match_id,
                    event.event_type.value,
                    event.minute,
                    event.extra_minute,
                    event.team_name,
                    event.player_name,
                    event.related_player_name,
                    event.period,
                    event.source_name,
                    event.source_event_id,
                    event.source_observation_id,
                ]
                for event in self._repository.list_events()
            ],
        )

        lineups_sheet = workbook.create_sheet("Lineups")
        _write_sheet(
            lineups_sheet,
            [
                "lineup_id",
                "match_id",
                "team",
                "player_name",
                "shirt_number",
                "position",
                "is_starter",
                "source_name",
            ],
            [
                [
                    row.lineup_id,
                    row.match_id,
                    row.team_name,
                    row.player_name,
                    row.shirt_number,
                    row.position,
                    row.is_starter,
                    row.source_name,
                ]
                for row in self._repository.list_lineups()
            ],
        )

        coverage_sheet = workbook.create_sheet("Coverage")
        _write_sheet(
            coverage_sheet,
            [
                "competition",
                "season",
                "field_name",
                "matches_total",
                "matches_present",
                "coverage_ratio",
                "missing_match_ids",
            ],
            [
                [
                    cell.competition_name,
                    cell.season_label,
                    cell.field_name,
                    cell.matches_total,
                    cell.matches_present,
                    cell.coverage_ratio,
                    ",".join(cell.missing_match_ids),
                ]
                for cell in CoverageReporter(self._repository).report()
            ],
        )

        sources_sheet = workbook.create_sheet("Sources")
        _write_sheet(
            sources_sheet,
            [
                "source_name",
                "source_match_id",
                "observation_id",
                "revision",
                "source_url",
                "retrieved_at",
                "source_timestamp",
                "raw_payload_hash",
                "confidence",
                "quality_flags",
            ],
            [
                [
                    record.source_name,
                    record.source_match_id,
                    record.observation_id,
                    record.revision,
                    record.source_url,
                    record.retrieved_at.isoformat(),
                    record.source_timestamp.isoformat() if record.source_timestamp else None,
                    record.raw_payload_hash,
                    record.confidence,
                    ",".join(flag.value for flag in record.quality_flags),
                ]
                for record in self._repository.list_source_records()
            ],
        )

        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        return destination


def read_sheet_rows(path: str | Path, sheet_name: str) -> list[list[object]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name]
    rows: list[list[object]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in row):
            continue
        rows.append(list(row))
    return rows


def _write_sheet(sheet: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(list(row))
