from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from sports_hedge.odds.coverage import CoverageReport
from sports_hedge.odds.models import MappingException, OddsObservation
from sports_hedge.odds.movement import OpenCloseMove

_HEADER_FONT = Font(bold=True)


def export_odds_workbook(
    path: str | Path,
    *,
    observations: list[OddsObservation],
    report: CoverageReport,
    moves: list[OpenCloseMove] | None = None,
    include_observations: bool = True,
    include_moves: bool = True,
) -> Path:
    """Write a review workbook. The SQLite repository remains source of truth."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    if include_observations:
        _write_observations(workbook.active, observations)
        _write_market_coverage(workbook.create_sheet("Market Coverage"), report)
    else:
        workbook.active.title = "Market Coverage"
        _write_market_coverage(workbook.active, report)
    _write_source_coverage(workbook.create_sheet("Source Coverage"), report)
    _write_exceptions(workbook.create_sheet("Mapping Exceptions"), report.mapping_exceptions)
    _write_quality(workbook.create_sheet("Quality Summary"), report)
    if include_moves:
        _write_open_close(workbook.create_sheet("Open Close Moves"), moves or [])
    workbook.save(destination)
    return destination


def _write_header(sheet: Worksheet, headers: list[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.font = _HEADER_FONT
    sheet.freeze_panes = "A2"


def _write_observations(sheet: Worksheet, observations: list[OddsObservation]) -> None:
    sheet.title = "Odds Observations"
    headers = [
        "observation_id",
        "canonical_match_id",
        "competition",
        "season",
        "home_team",
        "away_team",
        "kickoff_utc",
        "source",
        "venue",
        "bookmaker",
        "market_family",
        "period",
        "line",
        "selection",
        "side",
        "decimal_odds",
        "observed_at",
        "quote_type",
        "spread",
        "liquidity",
        "commission_known",
        "quality_tier",
        "venue_kind",
        "research_only",
        "implied_probability",
        "implied_logit",
        "semantics_complete",
        "settlement_key",
        "source_url",
        "retrieved_at",
        "raw_payload_hash",
    ]
    _write_header(sheet, headers)
    for row_index, observation in enumerate(observations, start=2):
        values = [
            observation.observation_id,
            observation.canonical_match_id,
            observation.competition_code,
            observation.season,
            observation.home_team,
            observation.away_team,
            observation.kickoff_utc.isoformat(),
            observation.source,
            observation.venue,
            observation.bookmaker,
            observation.market_family.value,
            observation.period.value,
            None if observation.line is None else format(observation.line, "f"),
            observation.selection,
            None if observation.side is None else observation.side.value,
            None if observation.decimal_odds is None else format(observation.decimal_odds, "f"),
            None if observation.observed_at is None else observation.observed_at.isoformat(),
            observation.quote_type.value,
            None if observation.spread is None else format(observation.spread, "f"),
            None if observation.liquidity is None else format(observation.liquidity, "f"),
            observation.commission_known,
            observation.quality_tier.value,
            observation.venue_kind.value,
            bool(observation.metadata.get("research_only")),
            None
            if observation.implied_probability() is None
            else format(observation.implied_probability(), "f"),
            observation.implied_logit(),
            observation.semantics_complete,
            observation.settlement_key,
            observation.source_url,
            observation.retrieved_at.isoformat(),
            observation.raw_payload_hash,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)


def _write_market_coverage(sheet: Worksheet, report: CoverageReport) -> None:
    headers = [
        "competition",
        "season",
        "market_family",
        "matches_in_repository",
        "expected_league_matches",
        "matches_with_market",
        "matches_with_opening_1x2",
        "matches_with_closing_1x2",
        "matches_with_timestamped",
        "matches_with_liquidity",
        "coverage_ratio",
        "missing_match_count",
        "calendar_gap",
    ]
    _write_header(sheet, headers)
    for row_index, row in enumerate(report.market_coverage, start=2):
        values = [
            row.competition_code,
            row.season,
            row.market_family,
            row.matches_in_repository,
            row.expected_league_matches,
            row.matches_with_market,
            row.matches_with_opening_1x2,
            row.matches_with_closing_1x2,
            row.matches_with_timestamped,
            row.matches_with_liquidity,
            row.coverage_ratio,
            row.missing_match_count,
            row.calendar_gap,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)


def _write_source_coverage(sheet: Worksheet, report: CoverageReport) -> None:
    headers = [
        "source",
        "required",
        "available",
        "observations",
        "matches",
        "timestamped_observations",
        "liquidity_observations",
        "quality_tiers",
        "notes",
    ]
    _write_header(sheet, headers)
    for row_index, row in enumerate(report.source_coverage, start=2):
        values = [
            row.source,
            row.required,
            row.available,
            row.observations,
            row.matches,
            row.timestamped_observations,
            row.liquidity_observations,
            ", ".join(f"{key}={value}" for key, value in sorted(row.quality_tiers.items())),
            row.notes,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)


def _write_exceptions(sheet: Worksheet, exceptions: list[MappingException]) -> None:
    headers = ["exception_id", "source", "source_reference", "field", "reason", "detail", "retrieved_at"]
    _write_header(sheet, headers)
    for row_index, exception in enumerate(exceptions, start=2):
        values = [
            exception.exception_id,
            exception.source,
            exception.source_reference,
            exception.field,
            exception.reason,
            exception.detail,
            exception.retrieved_at.isoformat(),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)


def _write_quality(sheet: Worksheet, report: CoverageReport) -> None:
    headers = ["quality_tier", "observations", "matches", "meaning"]
    _write_header(sheet, headers)
    for row_index, row in enumerate(report.quality_summary, start=2):
        values = [row.quality_tier, row.observations, row.matches, row.meaning]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)


def _write_open_close(sheet: Worksheet, moves: list[OpenCloseMove]) -> None:
    headers = [
        "canonical_match_id",
        "competition",
        "season",
        "bookmaker",
        "market_family",
        "selection",
        "opening_line",
        "opening_odds",
        "closing_odds",
        "implied_probability_delta",
        "implied_logit_delta",
        "research_only",
    ]
    _write_header(sheet, headers)
    for row_index, move in enumerate(moves, start=2):
        values = [
            move.canonical_match_id,
            move.competition_code,
            move.season,
            move.bookmaker,
            move.market_family,
            move.selection,
            None if move.line is None else format(move.line, "f"),
            format(move.opening_odds, "f"),
            format(move.closing_odds, "f"),
            format(move.implied_probability_delta, "f"),
            move.implied_logit_delta,
            move.research_only,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
