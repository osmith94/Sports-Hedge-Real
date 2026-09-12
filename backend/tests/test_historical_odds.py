from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from sports_hedge.domain.football import (
    FootballPeriod,
    MarketFamily,
    SettlementFingerprint,
    SettlementScope,
)
from sports_hedge.domain.models import MarketSide
from sports_hedge.facts.catalog import CompetitionCode
from sports_hedge.facts.identity import (
    CanonicalMatchRef,
    NaiveKickoffError,
    build_match_ref,
    canonical_match_id,
)
from sports_hedge.odds.adapters.football_data import FootballDataCsvAdapter
from sports_hedge.odds.adapters.smarkets import SmarketsHistoricalAdapter, smarkets_limitations
from sports_hedge.odds.adapters.synthetic import SyntheticOddsAdapter
from sports_hedge.odds.coverage import build_coverage_report
from sports_hedge.odds.excel import export_odds_workbook
from sports_hedge.odds.ingestion import OddsIngestionService
from sports_hedge.odds.mapping import OddsMappingError, map_raw_record
from sports_hedge.odds.models import QualityTier, QuoteType, RawOddsRecord, VenueKind
from sports_hedge.odds.repository import SqliteOddsRepository

FOOTBALL_DATA_SAMPLE = """Div,Date,Time,HomeTeam,AwayTeam,FTHG,FTAG,B365H,B365D,B365A,B365CH,B365CD,B365CA,B365>2.5,B365<2.5,B365C>2.5,B365C<2.5
E0,16/08/2025,15:00,Arsenal,Liverpool,2,1,2.10,3.40,3.50,2.05,3.50,3.60,1.85,2.00,1.90,1.95
E0,17/08/2025,16:30,Chelsea,Man City,1,1,2.50,3.30,2.80,2.55,3.25,2.75,1.90,1.90,1.88,1.92
"""

UNKNOWN_COMPETITION_CSV = """Div,Date,Time,HomeTeam,AwayTeam,FTHG,FTAG,B365H,B365D,B365A
XX,16/08/2025,15:00,Alpha,Beta,0,0,2.00,3.00,4.00
"""


def _ingest_synthetic() -> tuple[SqliteOddsRepository, OddsIngestionService]:
    repository = SqliteOddsRepository()
    service = OddsIngestionService(repository)
    service.ingest_adapter(SyntheticOddsAdapter())
    return repository, service


def test_odds_uses_facts_match_sha256_identity() -> None:
    kickoff = datetime(2025, 8, 16, 17, 30, tzinfo=UTC)
    expected = canonical_match_id(
        competition_code="premier_league",
        season="2025/26",
        home_team="Arsenal",
        away_team="Chelsea",
        kickoff_utc=kickoff,
    )
    # Same fixture and ID as PR #35's published Arsenal–Chelsea example.
    assert expected == "match:f295bd6ca68b6926073e179d"
    mapped = map_raw_record(
        RawOddsRecord(
            source="synthetic",
            source_market_id="id-check",
            source_reference="id-check",
            competition="Premier League",
            home_team="Arsenal",
            away_team="Chelsea",
            kickoff_utc=kickoff,
            market_family=MarketFamily.MATCH_RESULT,
            period=FootballPeriod.FULL_TIME,
            selection="home",
            decimal_odds=Decimal("2.10"),
            quote_type=QuoteType.CLOSING,
            retrieved_at=datetime(2026, 9, 11, tzinfo=UTC),
            settlement=SettlementFingerprint(
                scope=SettlementScope.REGULATION_TIME,
                period=FootballPeriod.FULL_TIME,
                extra_time_included=False,
                penalties_included=False,
            ),
            semantics_complete=True,
        )
    )
    assert mapped.canonical_match_id == expected
    assert mapped.competition_code == CompetitionCode.PREMIER_LEAGUE.value



def test_timestamped_and_closing_observations_persist_side_by_side() -> None:
    repository, _ = _ingest_synthetic()
    observations = [
        item
        for item in repository.list_observations()
        if item.home_team == "arsenal"
        and item.away_team == "liverpool"
        and item.market_family == MarketFamily.MATCH_RESULT
        and item.selection == "home"
        and item.decimal_odds is not None
    ]
    quote_types = {item.quote_type for item in observations}
    tiers = {item.quality_tier for item in observations}
    assert QuoteType.TIMESTAMPED in quote_types
    assert QuoteType.CLOSING in quote_types
    assert QualityTier.A in tiers
    assert QualityTier.C in tiers
    timestamped = next(
        item
        for item in observations
        if item.quote_type == QuoteType.TIMESTAMPED and item.liquidity is not None
    )
    closing = next(item for item in observations if item.quote_type == QuoteType.CLOSING)
    assert timestamped.observed_at is not None
    assert timestamped.liquidity is not None
    assert closing.observed_at is None
    assert closing.liquidity is None
    assert timestamped.observation_id != closing.observation_id


def test_mapping_is_deterministic_across_aliases() -> None:
    left = build_match_ref(
        competition="E0",
        home_team="Man City",
        away_team="Man United",
        kickoff_utc=datetime(2025, 8, 17, 15, 30, tzinfo=UTC),
    )
    right = build_match_ref(
        competition="Premier League",
        home_team="Manchester City",
        away_team="Manchester United",
        kickoff_utc=datetime(2025, 8, 17, 15, 30, tzinfo=UTC),
    )
    assert left.canonical_match_id == right.canonical_match_id
    assert left.competition_code == CompetitionCode.PREMIER_LEAGUE
    assert left.canonical_match_id.startswith("match:")
    assert isinstance(left, CanonicalMatchRef)


def test_incomplete_settlement_cannot_masquerade_as_equivalent() -> None:
    repository, _ = _ingest_synthetic()
    rows = [
        item
        for item in repository.list_observations()
        if item.source_market_id in {"syn-pl-ars-liv-1x2", "syn-pl-ars-liv-1x2-incomplete"}
        and item.selection == "home"
    ]
    complete = next(item for item in rows if item.semantics_complete)
    incomplete = next(item for item in rows if not item.semantics_complete)
    assert complete.market_equivalence_key() is not None
    assert incomplete.market_equivalence_key() is None
    assert complete.market_equivalence_key() != incomplete.market_equivalence_key()


def test_coverage_exposes_missing_markets_and_calendar_gaps() -> None:
    repository, _ = _ingest_synthetic()
    report = build_coverage_report(
        matches=repository.list_matches(),
        observations=repository.list_observations(),
        exceptions=repository.list_exceptions(),
        smarkets_limitations=smarkets_limitations(),
    )
    assert report.unresolved_mapping_count == 0
    corners = next(
        row
        for row in report.market_coverage
        if row.competition_code == "premier_league" and row.market_family == "corners"
    )
    assert corners.matches_in_repository == 2
    assert corners.matches_with_market == 0
    assert corners.missing_match_count == 2
    assert corners.calendar_gap == 378
    championship_1x2 = next(
        row
        for row in report.market_coverage
        if row.competition_code == "championship" and row.market_family == "match_result"
    )
    assert championship_1x2.matches_in_repository == 1
    assert championship_1x2.matches_with_market == 0
    smarkets = next(row for row in report.source_coverage if row.source == "smarkets")
    assert smarkets.required is False
    assert smarkets.available is False
    quality = {row.quality_tier: row.observations for row in report.quality_summary}
    assert quality["A"] > 0
    assert quality["C"] > 0
    assert quality["D"] > 0


def test_provenance_and_quality_tiers_are_retained() -> None:
    repository, _ = _ingest_synthetic()
    la_liga = next(
        item
        for item in repository.list_observations()
        if item.competition_code == "la_liga" and item.selection == "home"
    )
    assert la_liga.source == "synthetic"
    assert la_liga.quality_tier == QualityTier.A
    assert la_liga.commission_known is True
    assert la_liga.raw_payload_hash
    championship = next(
        item
        for item in repository.list_observations()
        if item.competition_code == "championship"
    )
    assert championship.quality_tier == QualityTier.D
    assert championship.decimal_odds is None


def test_ingest_is_idempotent_and_restartable() -> None:
    repository, service = _ingest_synthetic()
    first = service.ingest_adapter(SyntheticOddsAdapter())
    assert first.observations_created == 0
    assert first.observations_duplicate > 0
    assert first.exceptions == 0
    assert repository.get_checkpoint("synthetic") == f"records:{first.records_seen}"


def test_ambiguous_mapping_fails_closed() -> None:
    repository = SqliteOddsRepository()
    service = OddsIngestionService(repository)
    record = RawOddsRecord(
        source="synthetic",
        source_reference="bad-comp",
        competition="Mystery League",
        home_team="Alpha",
        away_team="Beta",
        kickoff_utc=datetime(2025, 8, 16, 14, 0, tzinfo=UTC),
        market_family=MarketFamily.MATCH_RESULT,
        period=FootballPeriod.FULL_TIME,
        selection="home",
        decimal_odds=Decimal("2.00"),
        retrieved_at=datetime(2026, 9, 11, tzinfo=UTC),
    )
    result = service.ingest_records("synthetic", [record])
    assert result.exceptions == 1
    assert result.observations_created == 0
    assert repository.list_exceptions()[0].reason == "unknown_competition"


def test_identical_home_away_is_rejected() -> None:
    record = RawOddsRecord(
        source="synthetic",
        source_reference="same-teams",
        competition="Premier League",
        home_team="Arsenal",
        away_team="Arsenal",
        kickoff_utc=datetime(2025, 8, 16, 14, 0, tzinfo=UTC),
        market_family=MarketFamily.MATCH_RESULT,
        period=FootballPeriod.FULL_TIME,
        selection="home",
        decimal_odds=Decimal("2.00"),
        retrieved_at=datetime(2026, 9, 11, tzinfo=UTC),
    )
    try:
        map_raw_record(record)
        raise AssertionError("expected mapping failure")
    except OddsMappingError:
        pass


def test_football_data_csv_is_quality_c_without_invented_timestamps() -> None:
    retrieved = datetime(2026, 9, 11, 12, 0, tzinfo=UTC)
    repository = SqliteOddsRepository()
    service = OddsIngestionService(repository)
    service.ingest_adapter(FootballDataCsvAdapter(FOOTBALL_DATA_SAMPLE, retrieved_at=retrieved))
    odds = [item for item in repository.list_observations() if item.decimal_odds is not None]
    assert odds
    assert all(item.quality_tier == QualityTier.C for item in odds)
    assert all(item.observed_at is None for item in odds)
    assert all(item.liquidity is None for item in odds)
    assert {item.quote_type for item in odds} == {QuoteType.OPENING, QuoteType.CLOSING}
    unknown = FootballDataCsvAdapter(UNKNOWN_COMPETITION_CSV, retrieved_at=retrieved)
    rejected = OddsIngestionService(SqliteOddsRepository()).ingest_adapter(unknown)
    assert rejected.exceptions == 1
    assert rejected.observations_created == 0


def test_smarkets_is_optional_and_does_not_fetch_the_network() -> None:
    adapter = SmarketsHistoricalAdapter()
    assert adapter.required is False
    assert adapter.configured is False
    assert adapter.fetch() == []
    limitations = smarkets_limitations()
    assert limitations["required"] is False
    assert limitations["network_fetch_enabled"] is False
    assert limitations["public_api_historical_archive"] is False


def test_excel_export_has_required_sheets(tmp_path: Path) -> None:
    repository, _ = _ingest_synthetic()
    report = build_coverage_report(
        matches=repository.list_matches(),
        observations=repository.list_observations(),
        exceptions=repository.list_exceptions(),
        smarkets_limitations=smarkets_limitations(),
    )
    path = export_odds_workbook(
        tmp_path / "odds.xlsx",
        observations=repository.list_observations(),
        report=report,
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Odds Observations",
        "Market Coverage",
        "Source Coverage",
        "Mapping Exceptions",
        "Quality Summary",
    ]
    assert workbook["Odds Observations"].max_row > 1


def test_league_settlement_fingerprint_is_reused_when_complete() -> None:
    record = RawOddsRecord(
        source="synthetic",
        source_market_id="m1",
        source_reference="ok",
        competition="La Liga",
        home_team="Real Madrid",
        away_team="Barcelona",
        kickoff_utc=datetime(2025, 10, 26, 19, 0, tzinfo=UTC),
        market_family=MarketFamily.MATCH_RESULT,
        period=FootballPeriod.FULL_TIME,
        selection="home",
        side=MarketSide.BACK,
        decimal_odds=Decimal("2.40"),
        observed_at=datetime(2025, 10, 26, 17, 30, tzinfo=UTC),
        quote_type=QuoteType.TIMESTAMPED,
        liquidity=Decimal(100),
        venue_kind=VenueKind.EXCHANGE,
        retrieved_at=datetime(2026, 9, 11, tzinfo=UTC),
        settlement=SettlementFingerprint(
            scope=SettlementScope.REGULATION_TIME,
            period=FootballPeriod.FULL_TIME,
            extra_time_included=False,
            penalties_included=False,
        ),
        semantics_complete=True,
    )
    mapped = map_raw_record(record)
    assert mapped.settlement_key == record.settlement.deterministic_key()
    assert mapped.quality_tier == QualityTier.A


def test_facts_identity_rejects_naive_kickoff() -> None:
    naive_kickoff = datetime(2025, 8, 16, 17, 30, tzinfo=UTC).replace(tzinfo=None)
    with pytest.raises(NaiveKickoffError):
        canonical_match_id(
            competition_code="premier_league",
            season="2025/26",
            home_team="Arsenal",
            away_team="Chelsea",
            kickoff_utc=naive_kickoff,
        )


def test_naive_timestamps_are_rejected() -> None:
    naive_kickoff = datetime(2025, 8, 16, 14, 0, tzinfo=UTC).replace(tzinfo=None)
    with pytest.raises(ValidationError):
        RawOddsRecord(
            source="synthetic",
            source_reference="naive",
            competition="Premier League",
            home_team="Arsenal",
            away_team="Liverpool",
            kickoff_utc=naive_kickoff,
            market_family=MarketFamily.MATCH_RESULT,
            period=FootballPeriod.FULL_TIME,
            selection="home",
            decimal_odds=Decimal("2.00"),
            retrieved_at=datetime(2026, 9, 11, 12, 0, tzinfo=UTC),
        )


def test_corrected_odds_are_append_only() -> None:
    retrieved = datetime(2026, 9, 11, 12, 0, tzinfo=UTC)
    kickoff = datetime(2025, 8, 16, 14, 0, tzinfo=UTC)
    settlement = SettlementFingerprint(
        scope=SettlementScope.REGULATION_TIME,
        period=FootballPeriod.FULL_TIME,
        extra_time_included=False,
        penalties_included=False,
    )
    shared = {
        "source": "synthetic",
        "source_market_id": "syn-close-revision",
        "source_reference": "syn-close-revision",
        "competition": "Premier League",
        "home_team": "Arsenal",
        "away_team": "Liverpool",
        "kickoff_utc": kickoff,
        "market_family": MarketFamily.MATCH_RESULT,
        "period": FootballPeriod.FULL_TIME,
        "selection": "home",
        "quote_type": QuoteType.CLOSING,
        "retrieved_at": retrieved,
        "settlement": settlement,
        "semantics_complete": True,
    }
    original = RawOddsRecord(**shared, decimal_odds=Decimal("2.10"), raw_payload={"v": 1})
    correction = RawOddsRecord(**shared, decimal_odds=Decimal("2.20"), raw_payload={"v": 2})
    replay = RawOddsRecord(**shared, decimal_odds=Decimal("2.10"), raw_payload={"v": 1})
    repository = SqliteOddsRepository()
    service = OddsIngestionService(repository)
    assert service.ingest_records("synthetic", [original]).observations_created == 1
    assert service.ingest_records("synthetic", [correction]).observations_created == 1
    replay_result = service.ingest_records("synthetic", [replay])
    rows = [
        item
        for item in repository.list_observations()
        if item.source_market_id == "syn-close-revision"
    ]
    assert len(rows) == 2
    assert {item.decimal_odds for item in rows} == {Decimal("2.10"), Decimal("2.20")}
    assert len({item.source_observation_key for item in rows}) == 1
    assert replay_result.observations_created == 0
    assert replay_result.observations_duplicate == 1

